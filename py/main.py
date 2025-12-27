import os
import tempfile
from io import BytesIO
from typing import Any, List, Optional
from urllib.parse import urlparse

import httpx
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pydantic import BaseModel

app = FastAPI()

# CORS 설정
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # 개발 환경에서는 모든 origin 허용
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class ImageData(BaseModel):
    imagePrompt: str
    isPlaceholder: bool


class FabricObject(BaseModel):
    type: str
    left: float
    top: float
    originX: str = "center"
    originY: str = "center"
    width: Optional[float] = 100
    height: Optional[float] = 100
    fill: Optional[str] = "#000000"
    stroke: Optional[str] = ""
    strokeWidth: float = 0
    text: Optional[str] = ""
    fontSize: float = 16
    fontFamily: str = "Arial"
    fontWeight: str = "normal"
    textAlign: str = "left"
    opacity: float = 1.0
    angle: float = 0
    scaleX: float = 1.0
    scaleY: float = 1.0
    radius: float = 0
    # Line 객체용 좌표 (optional)
    x1: Optional[float] = None
    x2: Optional[float] = None
    y1: Optional[float] = None
    y2: Optional[float] = None
    # Image 객체용 필드 (optional)
    src: Optional[str] = None
    crossOrigin: Optional[str] = None
    data: Optional[ImageData] = None


class FabricSlide(BaseModel):
    version: str
    objects: List[FabricObject]
    background: str = "#ffffff"


class ConvertRequest(BaseModel):
    slides: List[FabricSlide]


def hex_to_rgb_alpha(hex_color: str) -> tuple:
    """Convert hex color to (RGB, alpha) tuple
    Returns: ((R, G, B), alpha_0_to_1)
    """
    hex_color = hex_color.lstrip("#")

    if len(hex_color) >= 6:
        rgb = tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))

        # Check if alpha channel is present (8 characters: #RRGGBBAA)
        if len(hex_color) == 8:
            alpha_hex = int(hex_color[6:8], 16)
            alpha = alpha_hex / 255.0  # Convert to 0-1 range
            return (rgb, alpha)
        else:
            return (rgb, 1.0)  # Fully opaque

    return ((0, 0, 0), 1.0)


def set_shape_opacity(shape, opacity: float):
    """Set shape fill opacity (0-1) using lxml manipulation

    Args:
        shape: PowerPoint shape object
        opacity: Opacity value (0=transparent, 1=opaque)
    """
    if opacity >= 1.0:
        return

    try:
        from lxml import etree

        # PowerPoint uses 0-100000 scale where 100000 = fully opaque
        alpha_value = int(opacity * 100000)

        # Access spPr (shape properties) -> solidFill -> srgbClr
        spPr = shape.element.spPr

        # Find solidFill element
        solid_fill = spPr.find(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill"
        )

        if solid_fill is not None:
            # Find srgbClr element
            srgb_clr = solid_fill.find(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"
            )

            if srgb_clr is not None:
                # Remove existing alpha elements if any
                for alpha_elem in srgb_clr.findall(
                    ".//{http://schemas.openxmlformats.org/drawingml/2006/main}alpha"
                ):
                    srgb_clr.remove(alpha_elem)

                # Add new alpha element
                alpha_elem = etree.Element(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha"
                )
                alpha_elem.set("val", str(alpha_value))
                srgb_clr.append(alpha_elem)

                print(f"✅ Set fill opacity to {opacity} (alpha={alpha_value})")

    except Exception as e:
        print(f"❌ Could not set fill opacity: {e}")
        import traceback

        traceback.print_exc()


def set_line_opacity(shape, opacity: float):
    """Set shape line/stroke opacity (0-1) using lxml manipulation

    Args:
        shape: PowerPoint shape object
        opacity: Opacity value (0=transparent, 1=opaque)
    """
    if opacity >= 1.0:
        return

    try:
        from lxml import etree

        # PowerPoint uses 0-100000 scale where 100000 = fully opaque
        alpha_value = int(opacity * 100000)

        # Access spPr (shape properties) -> ln (line) -> solidFill -> srgbClr
        spPr = shape.element.spPr

        # Find ln (line) element
        ln = spPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ln")

        if ln is not None:
            # Find solidFill within ln
            solid_fill = ln.find(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill"
            )

            if solid_fill is not None:
                # Find srgbClr element
                srgb_clr = solid_fill.find(
                    ".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"
                )

                if srgb_clr is not None:
                    # Remove existing alpha elements if any
                    for alpha_elem in srgb_clr.findall(
                        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}alpha"
                    ):
                        srgb_clr.remove(alpha_elem)

                    # Add new alpha element
                    alpha_elem = etree.Element(
                        "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha"
                    )
                    alpha_elem.set("val", str(alpha_value))
                    srgb_clr.append(alpha_elem)

                    print(f"✅ Set line opacity to {opacity} (alpha={alpha_value})")

    except Exception as e:
        print(f"❌ Could not set line opacity: {e}")
        import traceback

        traceback.print_exc()


def fabric_to_inches(value: float, is_width: bool = True) -> Inches:
    """Convert Fabric.js pixels to PowerPoint inches
    1280px = 10 inches (width)
    720px = 5.625 inches (height)
    """
    if is_width:
        return Inches(value * 10 / 1280)
    else:
        return Inches(value * 5.625 / 720)


async def download_image(url: str) -> BytesIO:
    """Download image from URL or load from local file and return as BytesIO object

    Args:
        url: Image URL or local path (e.g., /upload/images/xxx.png)

    Returns:
        BytesIO object containing image data
    """
    # Check if it's a local path (starts with /upload/)
    if url.startswith("/upload/"):
        # Local file path - convert to absolute path
        # Assuming the web server is running in the project root
        # /upload/images/xxx.png -> ../web/public/upload/images/xxx.png
        local_path = os.path.join(
            os.path.dirname(os.path.dirname(__file__)),  # Go up from py/ to project root
            "web",
            "public",
            url.lstrip("/"),
        )

        if os.path.exists(local_path):
            with open(local_path, "rb") as f:
                return BytesIO(f.read())
        else:
            raise FileNotFoundError(f"Local image file not found: {local_path}")

    # Remote URL - download via HTTP
    async with httpx.AsyncClient(timeout=30.0) as client:
        response = await client.get(url)
        response.raise_for_status()
        return BytesIO(response.content)


async def create_pptx_from_fabric(slides_data: List[FabricSlide]) -> str:
    """Convert Fabric.js JSON array to PowerPoint file"""
    prs = Presentation()
    prs.slide_width = Inches(10)  # 1280px -> 10 inches
    prs.slide_height = Inches(5.625)  # 720px -> 5.625 inches (16:9 ratio)

    for slide_data in slides_data:
        slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(slide_layout)

        # Set background color
        background = slide.background
        fill = background.fill
        fill.solid()
        rgb, alpha = hex_to_rgb_alpha(slide_data.background)
        fill.fore_color.rgb = RGBColor(*rgb)
        # Note: Background alpha is typically not supported in PowerPoint

        # Add objects
        for obj in slide_data.objects:
            try:
                # Handle Line type separately
                if obj.type == "Line":
                    # Line uses x1, y1, x2, y2 coordinates
                    if obj.x1 is None or obj.y1 is None or obj.x2 is None or obj.y2 is None:
                        print(f"⚠️  Skipping Line object with null coordinates")
                        continue

                    # Line 그리기 (connector 사용)
                    x1_inches = fabric_to_inches(obj.left + obj.x1, is_width=True)
                    y1_inches = fabric_to_inches(obj.top + obj.y1, is_width=False)
                    x2_inches = fabric_to_inches(obj.left + obj.x2, is_width=True)
                    y2_inches = fabric_to_inches(obj.top + obj.y2, is_width=False)

                    # PowerPoint에서는 connector를 사용하여 선 그리기
                    connector = slide.shapes.add_connector(
                        1,  # Straight connector
                        x1_inches,
                        y1_inches,
                        x2_inches,
                        y2_inches,
                    )

                    stroke_rgb, stroke_alpha = hex_to_rgb_alpha(obj.stroke)
                    connector.line.color.rgb = RGBColor(*stroke_rgb)
                    connector.line.width = Pt(obj.strokeWidth)
                    set_line_opacity(connector, stroke_alpha * obj.opacity)
                    continue

                # For non-Line objects, width and height are required
                if obj.width is None or obj.height is None:
                    print(f"⚠️  Skipping {obj.type} object with null dimensions")
                    continue

                # Calculate actual dimensions with scale
                actual_width = obj.width * obj.scaleX
                actual_height = obj.height * obj.scaleY

                # Calculate position (adjust for origin)
                # Fabric.js uses center origin by default in v7
                left = obj.left
                top = obj.top

                if obj.originX == "center":
                    left -= actual_width / 2
                elif obj.originX == "right":
                    left -= actual_width

                if obj.originY == "center":
                    top -= actual_height / 2
                elif obj.originY == "bottom":
                    top -= actual_height

                # Convert to inches
                left_inches = fabric_to_inches(left, is_width=True)
                top_inches = fabric_to_inches(top, is_width=False)
                width_inches = fabric_to_inches(actual_width, is_width=True)
                height_inches = fabric_to_inches(actual_height, is_width=False)

                if obj.type == "Textbox":
                    textbox = slide.shapes.add_textbox(
                        left_inches, top_inches, width_inches, height_inches
                    )
                    text_frame = textbox.text_frame
                    text_frame.word_wrap = True
                    text_frame.text = obj.text

                    # Adjust vertical alignment and margins
                    text_frame.margin_top = Inches(0)
                    text_frame.margin_bottom = Inches(0)
                    text_frame.margin_left = Inches(0)
                    text_frame.margin_right = Inches(0)

                    # Set text properties
                    for paragraph in text_frame.paragraphs:
                        # Convert font size from Fabric.js pixels to PowerPoint points
                        # Fabric.js uses pixels, PPT uses points (1 point ≈ 1.333 pixels)
                        # Scale factor: (10 inches / 1280 px) * 72 points/inch ≈ 0.5625
                        ppt_font_size = obj.fontSize * 0.5625
                        paragraph.font.size = Pt(ppt_font_size)
                        paragraph.font.name = obj.fontFamily
                        paragraph.font.bold = obj.fontWeight == "bold"

                        # Set line spacing to 1.0 (single spacing)
                        # Default is usually 1.15-1.2 in PowerPoint
                        paragraph.line_spacing = 1.0

                        # Text alignment
                        if obj.textAlign == "center":
                            paragraph.alignment = PP_ALIGN.CENTER
                        elif obj.textAlign == "right":
                            paragraph.alignment = PP_ALIGN.RIGHT
                        else:
                            paragraph.alignment = PP_ALIGN.LEFT

                        # Text color (with alpha channel support)
                        rgb, alpha = hex_to_rgb_alpha(obj.fill)
                        paragraph.font.color.rgb = RGBColor(*rgb)
                        # Note: Text color alpha is not fully supported in python-pptx

                elif obj.type == "Rect":
                    shape = slide.shapes.add_shape(
                        1,  # Rectangle shape type
                        left_inches,
                        top_inches,
                        width_inches,
                        height_inches,
                    )
                    # Fill (with alpha channel from color)
                    shape.fill.solid()
                    rgb, color_alpha = hex_to_rgb_alpha(obj.fill)
                    shape.fill.fore_color.rgb = RGBColor(*rgb)

                    # Apply combined opacity: color alpha * object opacity
                    combined_opacity = color_alpha * obj.opacity
                    set_shape_opacity(shape, combined_opacity)

                    # Stroke (with alpha channel)
                    if obj.stroke and obj.strokeWidth > 0:
                        stroke_rgb, stroke_alpha = hex_to_rgb_alpha(obj.stroke)
                        shape.line.color.rgb = RGBColor(*stroke_rgb)
                        shape.line.width = Pt(obj.strokeWidth)
                        # Apply stroke opacity
                        set_line_opacity(shape, stroke_alpha * obj.opacity)
                    else:
                        shape.line.fill.background()

                elif obj.type == "Circle":
                    # Use oval shape for circles
                    # radius is already in pixels
                    diameter = obj.radius * 2 * obj.scaleX
                    diameter_inches = fabric_to_inches(diameter, is_width=True)

                    shape = slide.shapes.add_shape(
                        9,  # Oval shape type
                        left_inches,
                        top_inches,
                        diameter_inches,
                        diameter_inches,
                    )
                    shape.fill.solid()
                    rgb, color_alpha = hex_to_rgb_alpha(obj.fill)
                    shape.fill.fore_color.rgb = RGBColor(*rgb)

                    # Apply combined opacity: color alpha * object opacity
                    combined_opacity = color_alpha * obj.opacity
                    set_shape_opacity(shape, combined_opacity)

                    if obj.stroke and obj.strokeWidth > 0:
                        stroke_rgb, stroke_alpha = hex_to_rgb_alpha(obj.stroke)
                        shape.line.color.rgb = RGBColor(*stroke_rgb)
                        shape.line.width = Pt(obj.strokeWidth)
                        # Apply stroke opacity
                        set_line_opacity(shape, stroke_alpha * obj.opacity)
                    else:
                        shape.line.fill.background()

                elif obj.type == "Image":
                    # Image 객체 처리
                    # Check if it's a placeholder image that wasn't generated
                    if obj.data and obj.data.isPlaceholder:
                        print(
                            f"⚠️  Skipping placeholder image (not generated): {obj.data.imagePrompt[:50]}..."
                        )
                        continue

                    if not obj.src:
                        print(f"⚠️  Skipping Image object without src")
                        continue

                    try:
                        # Download image from URL or load from local file
                        image_stream = await download_image(obj.src)

                        # Add image to slide
                        picture = slide.shapes.add_picture(
                            image_stream, left_inches, top_inches, width_inches, height_inches
                        )

                        # Apply rotation if needed
                        if obj.angle != 0:
                            picture.rotation = obj.angle

                        print(f"✅ Added image from {obj.src[:50]}...")

                    except Exception as img_error:
                        print(f"❌ Failed to add image from {obj.src}: {img_error}")
                        # Continue processing other objects even if image fails
                        continue

            except Exception as e:
                print(f"Error processing object {obj.type}: {e}")
                continue

    # Save to temporary file
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    prs.save(temp_file.name)
    temp_file.close()

    return temp_file.name


@app.get("/")
def read_root():
    return {"message": "Fabric PPT API Server"}


@app.post("/convert")
async def convert_to_pptx(request: ConvertRequest):
    """Convert Fabric.js JSON array to PowerPoint file"""
    try:
        if not request.slides:
            raise HTTPException(status_code=400, detail="No slides provided")

        pptx_path = await create_pptx_from_fabric(request.slides)

        return FileResponse(
            pptx_path,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            filename="presentation.pptx",
            background=None,
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8731)
